from fastapi import FastAPI, BackgroundTasks, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
import openai, json, os, re, tempfile, shutil, math, subprocess, time
from datetime import date
from lxml import etree
from typing import Optional


# ──────────────────────────────────────────────
# 앱 초기화
# ──────────────────────────────────────────────
app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"], allow_headers=["*"],
)

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")


# ──────────────────────────────────────────────
# 한자 후처리
# ──────────────────────────────────────────────
def filter_hanja(text: str) -> str:
    return re.sub(r'[\u4e00-\u9fff\u3400-\u4dbf\uf900-\ufaff]', '', text)

def sanitize_data(data):
    if isinstance(data, dict):
        return {k: sanitize_data(v) for k, v in data.items()}
    if isinstance(data, str):
        return filter_hanja(data)
    return data


# ──────────────────────────────────────────────
# Whisper 청크 분할 STT
# ──────────────────────────────────────────────
def transcribe_in_chunks(file_path, model, chunk_minutes=5):
    result = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration",
         "-of", "default=noprint_wrappers=1:nokey=1", file_path],
        capture_output=True, text=True
    )
    duration = float(result.stdout.strip())
    chunk_sec = chunk_minutes * 60
    chunks = math.ceil(duration / chunk_sec)

    full_text = ""
    for i in range(chunks):
        start = i * chunk_sec
        chunk_path = f"{file_path}_chunk{i}.m4a"
        subprocess.run([
            "ffmpeg", "-i", file_path,
            "-ss", str(start), "-t", str(chunk_sec),
            "-c", "copy", chunk_path, "-y"
        ], capture_output=True)
        chunk_result = model.transcribe(chunk_path, language="ko")
        full_text += chunk_result["text"] + " "
        os.remove(chunk_path)

    return full_text.strip()


# ──────────────────────────────────────────────
# Gemini 호출 함수
# ──────────────────────────────────────────────
def _gemini_call(prompt):
    from google import genai
    client = genai.Client(api_key=GEMINI_API_KEY)
    for attempt in range(3):
        try:
            res = client.models.generate_content(model="gemini-3.6-flash", contents=prompt)
            clean = re.sub(r"```json|```", "", res.text).strip()
            return json.loads(clean)
        except Exception as e:
            if attempt < 2:
                time.sleep(3)
            else:
                raise e

def call_gemini_minutes(transcript, title, members):
    prompt = f"""JSON만 반환. 한국어. 한자 절대 금지.
{{"agenda":[{{"title":"안건 제목","content":"안건 내용 요약"}}],"summary":"회의 전체 내용을 3줄로 요약","action_items":[{{"content":"해야 할 일","deadline":"기한 (없으면 미정')"}}],"next_agenda":"다음 회의에서 논의할 안건"}}

회의 제목: {title}
참석자: {members}

회의 내용:
{transcript}"""
    return _gemini_call(prompt)


# ──────────────────────────────────────────────
# Request 모델
# ──────────────────────────────────────────────
class MinutesDocRequest(BaseModel):
    title: str
    date: Optional[str] = ""
    members: Optional[str] = ""
    transcript: Optional[str] = ""
    agenda: Optional[list] = []
    summary: Optional[str] = ""
    action_items: Optional[list] = []
    next_agenda: Optional[str] = ""


# ──────────────────────────────────────────────
# 엔드포인트
# ──────────────────────────────────────────────
@app.get("/")
def root():
    return {"status": "ok", "service": "minutes-ai API"}

@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/analyze-minutes")
async def analyze_minutes(
    file: UploadFile = File(...),
    title: str = Form("회의"),
    date: str = Form(""),
    members: str = Form(""),
    model: str = Form("gpt"),
):
    tmp_dir = tempfile.mkdtemp()
    try:
        audio_path = os.path.join(tmp_dir, file.filename)
        with open(audio_path, "wb") as f:
            f.write(await file.read())

        client_stt = openai.OpenAI(api_key=OPENAI_API_KEY)
        with open(audio_path, "rb") as audio_file:
            transcript = client_stt.audio.transcriptions.create(
                model="whisper-1",
                file=audio_file,
                language="ko"
            ).text

        if model == "gemini":
            summary = call_gemini_minutes(transcript, title, members)
            summary = sanitize_data(summary)
            return {"title": title, "date": date, "members": members, "transcript": transcript, **summary}

        client = openai.OpenAI(api_key=OPENAI_API_KEY)
        res = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": """당신은 회의록 작성 전문가입니다. 회의 내용을 분석하여 아래 JSON 형식으로 반환하세요.

규칙:
1. JSON만 반환, 다른 텍스트 절대 금지
2. 모든 내용은 한국어로 작성
3. 한자 사용 절대 금지

{
  "agenda": [
    {"title": "안건 제목", "content": "안건 내용 요약"}
  ],
  "summary": "회의 전체 내용을 3줄로 요약",
  "action_items": [
    {"content": "해야 할 일", "deadline": "기한 (없으면 '미정')"}
  ],
  "next_agenda": "다음 회의에서 논의할 안건"
}"""},
                {"role": "user", "content": f"회의 제목: {title}\n참석자: {members}\n\n회의 내용:\n{transcript}"}
            ]
        )
        clean = re.sub(r"```json|```", "", res.choices[0].message.content).strip()
        summary = json.loads(clean)
        summary = sanitize_data(summary)

        return {
            "title": title,
            "date": date,
            "members": members,
            "transcript": transcript,
            **summary
        }

    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


@app.post("/generate-minutes-word")
def generate_minutes_word(req: MinutesDocRequest, background_tasks: BackgroundTasks):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    today = req.date or date.today().strftime("%Y.%m.%d")
    doc_num = f"MOD-{today.replace('.','').replace('-','')[:8]}-001"

    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(3.0)

    COLOR_PRIMARY = RGBColor(0x1B, 0x3A, 0x6B)
    COLOR_ACCENT  = RGBColor(0x2E, 0x5F, 0xA3)
    COLOR_GRAY    = RGBColor(0x66, 0x66, 0x66)
    FONT_NAME     = "맑은 고딕"

    def add_run(para, text, bold=False, size=10, color=None):
        run = para.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        return run

    def set_cell_bg(cell, hex_color):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

    def add_section_title(label):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(28)
        p.paragraph_format.space_after  = Pt(8)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '18')
        left.set(qn('w:space'), '8')
        left.set(qn('w:color'), '2E5FA3')
        pBdr.append(left)
        pPr.append(pBdr)
        p.paragraph_format.left_indent = Cm(0.4)
        add_run(p, label, bold=True, size=12, color=COLOR_PRIMARY)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    add_run(title_p, "회  의  록", bold=True, size=20, color=COLOR_PRIMARY)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(2)
    add_run(sub_p, req.title, bold=True, size=12, color=COLOR_ACCENT)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.space_after = Pt(14)
    add_run(date_p, f"{today}  |  modui.ai", size=9, color=COLOR_GRAY)

    doc.add_paragraph()

    add_section_title("1. 회의 개요")
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = 'Table Grid'
    labels = ["문서번호", "문서명", "회의명", "일  시", "참 석 자"]
    values = [doc_num, "회의록", req.title, today, req.members]
    for i, (lbl, val) in enumerate(zip(labels, values)):
        row = info_table.rows[i]
        row.cells[0].width = Cm(3.0)
        row.cells[1].width = Cm(12.0)
        set_cell_bg(row.cells[0], 'E8EEF7')
        lp = row.cells[0].paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(lp, lbl, bold=True, size=9, color=COLOR_PRIMARY)
        vp = row.cells[1].paragraphs[0]
        add_run(vp, val or "-", size=9)
    doc.add_paragraph()
    doc.add_paragraph()

    add_section_title("2. 안건 및 논의 내용")
    ag_table = doc.add_table(rows=1 + len(req.agenda), cols=3)
    ag_table.style = 'Table Grid'
    hrow = ag_table.rows[0]
    for cell, label, w in zip(hrow.cells, ["번호","안건","내용"], [Cm(1.5),Cm(4.0),Cm(9.5)]):
        cell.width = w
        set_cell_bg(cell, '1B3A6B')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, label, bold=True, size=9, color=RGBColor(0xFF,0xFF,0xFF))
    for i, item in enumerate(req.agenda or []):
        row = ag_table.rows[i+1]
        bg = 'FFFFFF' if i % 2 == 0 else 'E8EEF7'
        for cell in row.cells:
            set_cell_bg(cell, bg)
        row.cells[0].width = Cm(1.5)
        row.cells[1].width = Cm(4.0)
        row.cells[2].width = Cm(9.5)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p0, str(i+1), bold=True, size=9, color=COLOR_ACCENT)
        add_run(row.cells[1].paragraphs[0], item.get("title",""), bold=True, size=9)
        add_run(row.cells[2].paragraphs[0], item.get("content",""), size=9)
    doc.add_paragraph()
    doc.add_paragraph()

    add_section_title("3. 회의 요약")
    for line in (req.summary or "").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(3)
        add_run(p, "• ", bold=True, color=COLOR_ACCENT)
        add_run(p, line.strip(), size=9)
    doc.add_paragraph()
    doc.add_paragraph()

    add_section_title("4. 결정사항 및 후속 조치")
    ai_table = doc.add_table(rows=1 + len(req.action_items), cols=3)
    ai_table.style = 'Table Grid'
    hrow2 = ai_table.rows[0]
    for cell, label, w in zip(hrow2.cells, ["No.","할 일","기한"], [Cm(1.2),Cm(10.3),Cm(3.5)]):
        cell.width = w
        set_cell_bg(cell, '2E5FA3')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, label, bold=True, size=9, color=RGBColor(0xFF,0xFF,0xFF))
    for i, item in enumerate(req.action_items or []):
        row = ai_table.rows[i+1]
        bg = 'FFFFFF' if i % 2 == 0 else 'E8EEF7'
        for cell in row.cells:
            set_cell_bg(cell, bg)
        row.cells[0].width = Cm(1.2)
        row.cells[1].width = Cm(10.3)
        row.cells[2].width = Cm(3.5)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p0, str(i+1), bold=True, size=9, color=COLOR_ACCENT)
        add_run(row.cells[1].paragraphs[0], item.get("content",""), size=9)
        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        deadline = item.get("deadline","미정")
        add_run(p2, deadline, bold=(deadline != "미정"), size=9,
                color=COLOR_GRAY if deadline == "미정" else COLOR_PRIMARY)
    doc.add_paragraph()

    add_section_title("5. 차기 회의 안건")
    next_p = doc.add_paragraph()
    next_p.paragraph_format.left_indent = Cm(0.4)
    add_run(next_p, req.next_agenda or "-", size=9)

    tmp_dir  = tempfile.mkdtemp()
    filename = f"{req.title}_회의록.docx"
    out_path = os.path.join(tmp_dir, filename)
    doc.save(out_path)

    background_tasks.add_task(shutil.rmtree, tmp_dir, ignore_errors=True)
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename,
        background=background_tasks,
    )