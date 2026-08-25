from fastapi import FastAPI, BackgroundTasks, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
import openai, json, zipfile, shutil, os, re, tempfile, math, subprocess, time
from datetime import date
from lxml import etree
from typing import Optional


# ──────────────────────────────────────────────
# 앱 초기화
# ──────────────────────────────────────────────
app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"], allow_headers=["*"],
)

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
TEMPLATE_PATH  = os.path.join(os.path.dirname(__file__), "template.pptx")

NS      = "http://schemas.openxmlformats.org/drawingml/2006/main"
PNS     = "http://schemas.openxmlformats.org/presentationml/2006/main"
RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


# ──────────────────────────────────────────────
# 슬라이드 콤보
# ──────────────────────────────────────────────
COMBO_3  = [("slide1.xml","cover"),("slide4.xml","cards"),("slide15.xml","outro")]
COMBO_5  = [("slide1.xml","cover"),("slide2.xml","overview"),("slide4.xml","cards"),("slide8.xml","analysis"),("slide15.xml","outro")]
COMBO_7  = [("slide1.xml","cover"),("slide2.xml","overview"),("slide4.xml","cards"),("slide7.xml","list"),("slide8.xml","analysis"),("slide10.xml","cards4"),("slide15.xml","outro")]
COMBO_10 = [("slide1.xml","cover"),("slide2.xml","overview"),("slide4.xml","cards"),("slide6.xml","keywords"),("slide7.xml","list"),("slide8.xml","analysis"),("slide10.xml","cards4"),("slide11.xml","timeline"),("slide2.xml","overview2"),("slide15.xml","outro")]
COMBOS   = {3:COMBO_3, 5:COMBO_5, 7:COMBO_7, 10:COMBO_10}


# ──────────────────────────────────────────────
# 한자 후처리
# ──────────────────────────────────────────────
def filter_hanja(text: str) -> str:
    return re.sub(r'[\u4e00-\u9fff\u3400-\u4dbf\uf900-\ufaff]', '', text)

def sanitize_data(data):
    if isinstance(data, dict):
        return {k: sanitize_data(v) for k, v in data.items()}
    if isinstance(data, str):
        return filter_hanja(data)
    return data


# ──────────────────────────────────────────────
# Whisper 청크 분할 STT
# ──────────────────────────────────────────────
def transcribe_in_chunks(file_path, model, chunk_minutes=5):
    result = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration",
         "-of", "default=noprint_wrappers=1:nokey=1", file_path],
        capture_output=True, text=True
    )
    duration = float(result.stdout.strip())
    chunk_sec = chunk_minutes * 60
    chunks = math.ceil(duration / chunk_sec)

    full_text = ""
    for i in range(chunks):
        start = i * chunk_sec
        chunk_path = f"{file_path}_chunk{i}.m4a"
        subprocess.run([
            "ffmpeg", "-i", file_path,
            "-ss", str(start), "-t", str(chunk_sec),
            "-c", "copy", chunk_path, "-y"
        ], capture_output=True)
        chunk_result = model.transcribe(chunk_path, language="ko")
        full_text += chunk_result["text"] + " "
        os.remove(chunk_path)

    return full_text.strip()


# ──────────────────────────────────────────────
# GPT 호출 - PPT
# ──────────────────────────────────────────────
class PPTRequest(BaseModel):
    keyword: str
    team: Optional[str] = "딸깍"
    slide_count: Optional[int] = 5
    model: Optional[str] = "gpt"

def call_gpt(keyword):
    client = openai.OpenAI(api_key=OPENAI_API_KEY)
    res = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {"role": "system", "content": """당신은 프레젠테이션 콘텐츠 작성 전문가입니다. 키워드를 받아 JSON을 생성하세요.
반드시 아래 규칙을 따르세요:
1. JSON만 반환, 다른 텍스트 절대 금지
2. 글자수 제한은 반드시 지켜야 하는 최대값입니다. 초과하면 절대 안됩니다
3. 본문 필드는 지정된 글자수를 최대한 채워서 작성
4. 제목은 간결하고 핵심만 담아야 합니다
5. 한자 사용 절대 금지. 强, 化, 性 등 한자 포함하면 안됨. 순한글 또는 영문만 사용
6. 카드 제목은 핵심 키워드로 작성, 숫자 나열(이점1, 항목2 등) 절대 금지

{
  "cover": {
    "title_line1": "제목 앞부분 (6자이내)",
    "title_line2": "제목 뒷부분 (8자이내)",
    "description": "발표 개요 (40자이내)",
    "subtitle": "부제목 (20자이내)"
  },
  "overview": {
    "title": "슬라이드 제목 (15자이내)",
    "section_title": "핵심 주제 문구 (15자이내)",
    "left_body": "왼쪽 본문. 3~4문장. 구체적 수치나 사례 포함. 반드시 180자 이상 200자 이내",
    "right_body": "오른쪽 본문. 3~4문장. 구체적 내용 포함. 반드시 180자 이상 200자 이내"
  },
  "overview2": {
    "title": "두번째 슬라이드 제목 (15자이내)",
    "section_title": "두번째 핵심 주제 (15자이내)",
    "left_body": "왼쪽 본문. 3~4문장. 반드시 180자 이상 200자 이내",
    "right_body": "오른쪽 본문. 3~4문장. 반드시 180자 이상 200자 이내"
  },
  "cards": {
    "title": "카드 슬라이드 제목 (15자이내)",
    "card1_title": "카드1 키워드 (7자이내)",
    "card1_body": "카드1 내용. 2~3문장. 75자이상 90자이내",
    "card2_title": "카드2 키워드 (7자이내)",
    "card2_body": "카드2 내용. 2~3문장. 75자이상 90자이내",
    "card3_title": "카드3 키워드 (7자이내)",
    "card3_body": "카드3 내용. 2~3문장. 75자이상 90자이내"
  },
  "keywords": {
    "title": "키워드 슬라이드 제목 (15자이내)",
    "label1": "라벨1 (5자이내)",
    "label2": "라벨2 (5자이내)",
    "label3": "라벨3 (5자이내)",
    "label4": "라벨4 (5자이내)",
    "keyword1": "키워드1 (5자이내)",
    "keyword2": "키워드2 (5자이내)",
    "keyword3": "키워드3 (5자이내)",
    "keyword4": "키워드4 (5자이내)",
    "summary": "핵심 요약 문장 (50자이내)"
  },
  "list": {
    "title": "리스트 슬라이드 제목 (15자이내)",
    "intro": "리스트 소개 문장 (45자이내)",
    "item1": "항목1 (25자이내)",
    "item2": "항목2 (25자이내)",
    "item3": "항목3 (25자이내)"
  },
  "analysis": {
    "title": "분석 슬라이드 제목 (15자이내)",
    "cause1_title": "원인1 (3자이내)",
    "cause1_body": "원인1 설명. 마침표로 끝나는 완결된 1문장. 40자이내",
    "cause2_title": "원인2 (3자이내)",
    "cause2_body": "원인2 설명. 마침표로 끝나는 완결된 1문장. 40자이내",
    "cause3_title": "원인3 (3자이내)",
    "cause3_body": "원인3 설명. 마침표로 끝나는 완결된 1문장. 40자이내",
    "result": "결과 키워드 (4자이내)",
    "result_body": "결과 설명. 각 문장은 반드시 마침표로 끝낼 것. 4~5문장. 150자이내"
  },
  "cards4": {
    "title": "4카드 슬라이드 제목 (15자이내)",
    "card1_title": "카드1 키워드. 반드시 5자이내. 절대 초과금지",
    "card2_title": "카드2 키워드. 반드시 5자이내. 절대 초과금지",
    "card3_title": "카드3 키워드. 반드시 5자이내. 절대 초과금지",
    "card4_title": "카드4 키워드. 반드시 5자이내. 절대 초과금지",
    "card1_body": "카드1 설명. 마침표로 끝나는 완결된 1~2문장. 50자이내",
    "card2_body": "카드2 설명. 마침표로 끝나는 완결된 1~2문장. 50자이내",
    "card3_body": "카드3 설명. 마침표로 끝나는 완결된 1~2문장. 50자이내",
    "card4_body": "카드4 설명. 마침표로 끝나는 완결된 1~2문장. 50자이내"
  },
  "timeline": {
    "title": "타임라인 슬라이드 제목 (15자이내)",
    "step1_title": "제목 키워드. 반드시 5자이내. 절대 초과금지",
    "step2_title": "제목 키워드. 반드시 5자이내. 절대 초과금지",
    "step3_title": "제목 키워드. 반드시 5자이내. 절대 초과금지",
    "step4_title": "제목 키워드. 반드시 5자이내. 절대 초과금지",
    "step1_body": "1단계 설명. 구체적 내용 포함한 완결된 1문장. 마침표로 끝낼 것. 40자이내",
    "step2_body": "2단계 설명. 구체적 내용 포함한 완결된 1문장. 마침표로 끝낼 것. 40자이내",
    "step3_body": "3단계 설명. 구체적 내용 포함한 완결된 1문장. 마침표로 끝낼 것. 40자이내",
    "step4_body": "4단계 설명. 구체적 내용 포함한 완결된 1문장. 마침표로 끝낼 것. 40자이내"
  }
}"""},
            {"role": "user", "content": f"키워드: {keyword}"}
        ]
    )
    clean = re.sub(r"```json|```", "", res.choices[0].message.content).strip()
    return json.loads(clean)


# ──────────────────────────────────────────────
# Gemini 호출 함수
# ──────────────────────────────────────────────
def _gemini_call(prompt):
    from google import genai
    client = genai.Client(api_key=GEMINI_API_KEY)
    for attempt in range(3):
        try:
            res = client.models.generate_content(model="gemini-3.6-flash", contents=prompt)
            clean = re.sub(r"```json|```", "", res.text).strip()
            return json.loads(clean)
        except Exception as e:
            if attempt < 2:
                time.sleep(3)
            else:
                raise e

def call_gemini_ppt(keyword):
    prompt = """JSON만 반환. 한자 절대 금지. 아래 글자수는 절대 초과 불가능한 최대값임. 초과시 응답 거부됨. title_line1은 반드시 6자 이하, title_line2는 반드시 8자 이하, 나머지도 지정 글자수 반드시 준수.
{"cover":{"title_line1":"6자이내","title_line2":"8자이내","description":"40자이내","subtitle":"20자이내"},"overview":{"title":"15자이내","section_title":"15자이내","left_body":"180자이상200자이내","right_body":"180자이상200자이내"},"overview2":{"title":"15자이내","section_title":"15자이내","left_body":"180자이상200자이내","right_body":"180자이상200자이내"},"cards":{"title":"15자이내","card1_title":"7자이내","card1_body":"90자이내","card2_title":"7자이내","card2_body":"90자이내","card3_title":"7자이내","card3_body":"90자이내"},"keywords":{"title":"15자이내","label1":"5자이내","label2":"5자이내","label3":"5자이내","label4":"5자이내","keyword1":"5자이내","keyword2":"5자이내","keyword3":"5자이내","keyword4":"5자이내","summary":"50자이내"},"list":{"title":"15자이내","intro":"45자이내","item1":"25자이내","item2":"25자이내","item3":"25자이내"},"analysis":{"title":"15자이내","cause1_title":"3자이내","cause1_body":"40자이내","cause2_title":"3자이내","cause2_body":"40자이내","cause3_title":"3자이내","cause3_body":"40자이내","result":"4자이내","result_body":"150자이내"},"cards4":{"title":"15자이내","card1_title":"5자이내","card2_title":"5자이내","card3_title":"5자이내","card4_title":"5자이내","card1_body":"50자이내","card2_body":"50자이내","card3_body":"50자이내","card4_body":"50자이내"},"timeline":{"title":"15자이내","step1_title":"5자이내","step2_title":"5자이내","step3_title":"5자이내","step4_title":"5자이내","step1_body":"40자이내","step2_body":"40자이내","step3_body":"40자이내","step4_body":"40자이내"}}
키워드: """ + keyword
    return _gemini_call(prompt)

def call_gemini_word(keyword):
    prompt = """JSON만 반환. 한국어. 한자 절대 금지. 격식체. 아래 글자수는 절대 초과 불가. title은 반드시 20자 이하, subtitle은 반드시 30자 이하. 초과시 응답 거부됨.
{"title":"20자이내","subtitle":"30자이내","overview":{"heading":"1. 개요","background":"150자이내","purpose":"100자이내"},"main":{"heading":"2. 핵심 내용","section1_title":"15자이내","section1_body":"150자이내","section2_title":"15자이내","section2_body":"150자이내","section3_title":"15자이내","section3_body":"150자이내"},"analysis":{"heading":"3. 분석","cause1":"100자이내","cause2":"100자이내","cause3":"100자이내","result":"150자이내"},"conclusion":{"heading":"4. 결론","summary":"150자이내","expected":"100자이내","one_line_summary":"50자이내"}}
키워드: """ + keyword
    return _gemini_call(prompt)

def call_gemini_minutes(transcript, title, members):
    prompt = f"""JSON만 반환. 한국어. 한자 절대 금지.
{{"agenda":[{{"title":"안건 제목","content":"안건 내용 요약"}}],"summary":"회의 전체 내용을 3줄로 요약","action_items":[{{"content":"해야 할 일","deadline":"기한 (없으면 미정')"}}],"next_agenda":"다음 회의에서 논의할 안건"}}

회의 제목: {title}
참석자: {members}

회의 내용:
{transcript}"""
    return _gemini_call(prompt)

def call_ai_ppt(keyword, model="gpt"):
    if model == "gemini":
        return call_gemini_ppt(keyword)
    return call_gpt(keyword)


# ──────────────────────────────────────────────
# XML 헬퍼
# ──────────────────────────────────────────────
def find_shape(root, name):
    for sp in root.iter(f"{{{PNS}}}sp"):
        nvSpPr = sp.find(f"{{{PNS}}}nvSpPr")
        if nvSpPr is not None:
            cNvPr = nvSpPr.find(f"{{{PNS}}}cNvPr")
            if cNvPr is not None and cNvPr.get('name') == name:
                return sp
    return None

def set_text(root, name, txt):
    import copy
    sp = find_shape(root, name)
    if sp is None:
        return

    txBody = sp.find(f"{{{NS}}}txBody")
    if txBody is None:
        txBody = sp.find(f"{{{PNS}}}txBody")
    if txBody is None:
        return

    paras = txBody.findall(f"{{{NS}}}p")
    if not paras:
        return

    first_para = paras[0]
    for p in paras[1:]:
        txBody.remove(p)

    runs = first_para.findall(f"{{{NS}}}r")
    saved_rPr = None
    for r in runs:
        rPr = r.find(f"{{{NS}}}rPr")
        if rPr is not None:
            saved_rPr = copy.deepcopy(rPr)
            break

    for r in runs:
        first_para.remove(r)
    endPr = first_para.find(f"{{{NS}}}endParaRPr")
    if endPr is not None:
        first_para.remove(endPr)

    new_r = etree.SubElement(first_para, f"{{{NS}}}r")
    if saved_rPr is not None:
        new_r.append(saved_rPr)
    new_t = etree.SubElement(new_r, f"{{{NS}}}t")
    new_t.text = txt.strip()

def s(d, k, lim):
    text = (d.get(k) or "")[:lim]
    if not text:
        return text
    if text[-1] in '.!?':
        return text
    for punct in ['습니다.', '입니다.', '됩니다.', '있습니다.', '합니다.', '니다.', '다.', '요.', '.', '!', '?']:
        idx = text.rfind(punct)
        if idx != -1:
            return text[:idx + len(punct)]
    return text

def st_s(d, k, lim):
    text = (d.get(k) or "")
    if len(text) <= lim:
        return text
    truncated = text[:lim]
    last_space = truncated.rfind(' ')
    if last_space > lim // 2:
        return truncated[:last_space]
    return truncated

def st_title(root, name, val):
    set_text(root, name, val + " ")

def _get_bodyPr(sp):
    txBody = sp.find(f"{{{NS}}}txBody")
    if txBody is None:
        txBody = sp.find(f"{{{PNS}}}txBody")
    if txBody is None:
        return None
    return txBody.find(f"{{{NS}}}bodyPr")

def set_no_autofit(root, name):
    for sp in root.iter(f"{{{PNS}}}sp"):
        nvSpPr = sp.find(f"{{{PNS}}}nvSpPr")
        if nvSpPr is None:
            continue
        cNvPr = nvSpPr.find(f"{{{PNS}}}cNvPr")
        if cNvPr is None or cNvPr.get("name") != name:
            continue
        bodyPr = _get_bodyPr(sp)
        if bodyPr is None:
            continue
        for tag in [f"{{{NS}}}normAutofit", f"{{{NS}}}spAutoFit"]:
            el = bodyPr.find(tag)
            if el is not None:
                bodyPr.remove(el)
        if bodyPr.find(f"{{{NS}}}noAutofit") is None:
            etree.SubElement(bodyPr, f"{{{NS}}}noAutofit")

def set_norm_autofit(root, name):
    for sp in root.iter(f"{{{PNS}}}sp"):
        nvSpPr = sp.find(f"{{{PNS}}}nvSpPr")
        if nvSpPr is None:
            continue
        cNvPr = nvSpPr.find(f"{{{PNS}}}cNvPr")
        if cNvPr is None or cNvPr.get("name") != name:
            continue
        bodyPr = _get_bodyPr(sp)
        if bodyPr is None:
            continue
        for tag in [f"{{{NS}}}noAutofit", f"{{{NS}}}spAutoFit"]:
            el = bodyPr.find(tag)
            if el is not None:
                bodyPr.remove(el)
        if bodyPr.find(f"{{{NS}}}normAutofit") is None:
            etree.SubElement(bodyPr, f"{{{NS}}}normAutofit")

def st(root, name, val):
    set_text(root, name, val)


# ──────────────────────────────────────────────
# 슬라이드별 fill 함수
# ──────────────────────────────────────────────
def fill_cover(root, data, keyword, team, today):
    c = data.get("cover", {})
    st_title(root, "TextBox 5",  st_s(c, "title_line1", 10))
    st_title(root, "TextBox 6",  st_s(c, "title_line2", 12))
    set_norm_autofit(root, "TextBox 5")
    set_norm_autofit(root, "TextBox 6")
    st(root, "TextBox 7",  s(c, "description", 40))
    st(root, "TextBox 8",  s(c, "subtitle", 20))
    st(root, "TextBox 9",  f"팀  {team}")
    st(root, "TextBox 10", today)
    st(root, "TextBox 11", "modui.ai")

def fill_overview(root, dk, data, label):
    o = data.get(dk, {})
    st_title(root, "TextBox 19", st_s(o, "title", 15))
    set_norm_autofit(root, "TextBox 19")
    st(root, "TextBox 20", "")
    st(root, "TextBox 21", label)
    st(root, "TextBox 22", s(o, "left_body", 200))
    st(root, "TextBox 23", s(o, "right_body", 200))
    st(root, "TextBox 24", st_s(o, "section_title", 15))

def fill_cards(root, dk, data, label):
    c = data.get(dk, {})
    st_title(root, "TextBox 17", st_s(c, "title", 15))
    set_norm_autofit(root, "TextBox 17")
    st(root, "TextBox 18", "")
    st(root, "TextBox 19", label)
    st(root, "TextBox 29", s(c, "card1_body", 95))
    st(root, "TextBox 30", s(c, "card2_body", 95))
    st(root, "TextBox 31", s(c, "card3_body", 95))
    st(root, "TextBox 32", st_s(c, "card1_title", 7))
    st(root, "TextBox 33", st_s(c, "card2_title", 7))
    st(root, "TextBox 34", st_s(c, "card3_title", 7))
    st(root, "TextBox 38", "Card 01")
    st(root, "TextBox 39", "Card 02")
    st(root, "TextBox 40", "Card 03")
    for box in ["TextBox 29", "TextBox 30", "TextBox 31",
                "TextBox 32", "TextBox 33", "TextBox 34"]:
        set_norm_autofit(root, box)

def fill_keywords(root, dk, data, label):
    k = data.get(dk, {})
    st_title(root, "TextBox 17", st_s(k, "title", 12))
    set_norm_autofit(root, "TextBox 17")
    st(root, "TextBox 18", "")
    st(root, "TextBox 19", label)
    st(root, "TextBox 39", s(k, "label1", 5))
    st(root, "TextBox 40", s(k, "label2", 5))
    st(root, "TextBox 41", s(k, "label3", 5))
    st(root, "TextBox 42", s(k, "label4", 5))
    st(root, "TextBox 43", s(k, "keyword1", 5))
    st(root, "TextBox 44", s(k, "keyword2", 5))
    st(root, "TextBox 45", s(k, "keyword3", 5))
    st(root, "TextBox 46", s(k, "keyword4", 5))
    st(root, "TextBox 51", s(k, "summary", 50))

def fill_list(root, dk, data, label):
    li = data.get(dk, {})
    st_title(root, "TextBox 17", st_s(li, "title", 15))
    set_norm_autofit(root, "TextBox 17")
    st(root, "TextBox 18", "")
    st(root, "TextBox 19", label)
    st(root, "TextBox 20", s(li, "intro", 45))
    st(root, "TextBox 30", s(li, "item1", 25))
    st(root, "TextBox 31", s(li, "item2", 25))
    st(root, "TextBox 32", s(li, "item3", 25))

def fill_analysis(root, dk, data, label):
    a = data.get(dk, {})
    st_title(root, "TextBox 17", st_s(a, "title", 15))
    set_norm_autofit(root, "TextBox 17")
    st(root, "TextBox 18", "")
    st(root, "TextBox 19", label)
    st(root, "TextBox 33", s(a, "cause1_body", 45))
    st(root, "TextBox 34", s(a, "cause2_body", 45))
    st(root, "TextBox 35", s(a, "cause3_body", 45))
    st(root, "TextBox 48", st_s(a, "cause1_title", 3))
    st(root, "TextBox 49", st_s(a, "cause2_title", 3))
    st(root, "TextBox 50", st_s(a, "cause3_title", 3))
    st(root, "TextBox 51", s(a, "result_body", 160))
    st(root, "TextBox 52", "분석 결과")
    for box in ["TextBox 33", "TextBox 34", "TextBox 35", "TextBox 51"]:
        set_norm_autofit(root, box)

def fill_cards4(root, dk, data, label):
    c = data.get(dk, {})
    st_title(root, "TextBox 17", st_s(c, "title", 15))
    set_norm_autofit(root, "TextBox 17")
    st(root, "TextBox 18", "")
    st(root, "TextBox 19", label)
    st(root, "TextBox 32", st_s(c, "card1_title", 5))
    st(root, "TextBox 33", st_s(c, "card2_title", 5))
    st(root, "TextBox 34", st_s(c, "card3_title", 5))
    st(root, "TextBox 35", st_s(c, "card4_title", 5))
    st(root, "TextBox 36", s(c, "card1_body", 55))
    st(root, "TextBox 37", s(c, "card2_body", 55))
    st(root, "TextBox 38", s(c, "card3_body", 55))
    st(root, "TextBox 39", s(c, "card4_body", 55))
    for box in ["TextBox 32", "TextBox 33", "TextBox 34", "TextBox 35",
                "TextBox 36", "TextBox 37", "TextBox 38", "TextBox 39"]:
        set_norm_autofit(root, box)

def fill_timeline(root, dk, data, label):
    t = data.get(dk, {})
    st_title(root, "TextBox 17", st_s(t, "title", 15))
    set_norm_autofit(root, "TextBox 17")
    st(root, "TextBox 18", "")
    st(root, "TextBox 19", label)
    st(root, "TextBox 25", st_s(t, "step1_title", 5))
    st(root, "TextBox 28", st_s(t, "step2_title", 5))
    st(root, "TextBox 30", st_s(t, "step3_title", 5))
    st(root, "TextBox 32", st_s(t, "step4_title", 5))
    st(root, "TextBox 26", s(t, "step1_body", 40))
    st(root, "TextBox 40", s(t, "step2_body", 40))
    st(root, "TextBox 31", s(t, "step3_body", 40))
    st(root, "TextBox 41", s(t, "step4_body", 40))
    for box in ["TextBox 17", "TextBox 25", "TextBox 28", "TextBox 30", "TextBox 32"]:
        set_no_autofit(root, box)
    for box in ["TextBox 26", "TextBox 40", "TextBox 31", "TextBox 41"]:
        set_norm_autofit(root, box)

def fill_outro(root, team, today):
    st(root, "TextBox 7",  "")
    st(root, "TextBox 8",  "")
    st(root, "TextBox 9",  f"팀  {team}")
    st(root, "TextBox 10", today)
    st(root, "TextBox 11", "modui.ai")

FILL = {
    "cover":     fill_cover,
    "overview":  fill_overview,
    "overview2": fill_overview,
    "cards":     fill_cards,
    "keywords":  fill_keywords,
    "list":      fill_list,
    "analysis":  fill_analysis,
    "cards4":    fill_cards4,
    "timeline":  fill_timeline,
    "outro":     fill_outro,
}


# ──────────────────────────────────────────────
# 슬라이드 복사 / 등록
# ──────────────────────────────────────────────
def copy_slide(unpacked, src, dst):
    sd = os.path.join(unpacked, "ppt", "slides")
    rd = os.path.join(sd, "_rels")
    shutil.copy(os.path.join(sd, src), os.path.join(sd, dst))
    sr = src + ".rels"
    dr = dst + ".rels"
    if os.path.exists(os.path.join(rd, sr)):
        shutil.copy(os.path.join(rd, sr), os.path.join(rd, dr))

def register_slide(fname, prs_root, rels_root):
    PRS2  = "http://schemas.openxmlformats.org/presentationml/2006/main"
    R2    = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    RELS2 = "http://schemas.openxmlformats.org/package/2006/relationships"
    max_rid = max(
        (int(r.get("Id", "rId0")[3:]) for r in rels_root.findall(f"{{{RELS2}}}Relationship")
         if r.get("Id", "").startswith("rId")),
        default=0
    )
    new_rid = f"rId{max_rid+1}"
    etree.SubElement(rels_root, f"{{{RELS2}}}Relationship",
                     {"Id": new_rid,
                      "Type": f"{R2}/slide",
                      "Target": f"slides/{fname}"})
    sld_lst = prs_root.find(f"{{{PRS2}}}sldIdLst")
    max_sid = max((int(s.get("id", "255")) for s in sld_lst), default=255)
    etree.SubElement(sld_lst, f"{{{PRS2}}}sldId",
                     {"id": str(max_sid + 1), f"{{{R2}}}id": new_rid})


# ──────────────────────────────────────────────
# PPT 생성 엔드포인트
# ──────────────────────────────────────────────
@app.post("/generate-ppt")
def generate_ppt(req: PPTRequest, background_tasks: BackgroundTasks):
    keyword     = req.keyword
    team        = req.team or "딸깍"
    slide_count = req.slide_count or 5
    today       = date.today().strftime("%Y.%m.%d")

    best  = min(COMBOS, key=lambda x: abs(x - slide_count))
    combo = COMBOS[best]

    data = call_gemini_ppt(keyword) if req.model == "gemini" else call_gpt(keyword)
    data = sanitize_data(data)
    print(json.dumps(data, ensure_ascii=False, indent=2))

    tmp_dir  = tempfile.mkdtemp()
    unpacked = os.path.join(tmp_dir, "unpacked")
    shutil.copy(TEMPLATE_PATH, os.path.join(tmp_dir, "output.pptx"))
    with zipfile.ZipFile(os.path.join(tmp_dir, "output.pptx"), "r") as z:
        z.extractall(unpacked)

    prs_path  = os.path.join(unpacked, "ppt", "presentation.xml")
    rels_path = os.path.join(unpacked, "ppt", "_rels", "presentation.xml.rels")
    with open(prs_path,  "rb") as f: prs_root  = etree.fromstring(f.read())
    with open(rels_path, "rb") as f: rels_root = etree.fromstring(f.read())

    PRS2  = "http://schemas.openxmlformats.org/presentationml/2006/main"
    R2    = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    RELS2 = "http://schemas.openxmlformats.org/package/2006/relationships"

    sld_lst = prs_root.find(f"{{{PRS2}}}sldIdLst")
    for e in list(sld_lst):
        sld_lst.remove(e)
    for r in list(rels_root.findall(f"{{{RELS2}}}Relationship")):
        t = r.get("Type", "")
        if "/slide" in t and "/slideLayout" not in t and "/slideMaster" not in t:
            rels_root.remove(r)

    used = {}
    num  = [0]
    def lbl():
        num[0] += 1
        return f"{num[0]:02d}"

    for tmpl, dk in combo:
        used[tmpl] = used.get(tmpl, 0) + 1
        actual = tmpl if used[tmpl] == 1 else tmpl.replace(".xml", f"_copy{used[tmpl]}.xml")
        if used[tmpl] > 1:
            copy_slide(unpacked, tmpl, actual)
        register_slide(actual, prs_root, rels_root)

        path = os.path.join(unpacked, "ppt", "slides", actual)
        with open(path, "rb") as f:
            xml = f.read()
        root = etree.fromstring(xml)

        fn = FILL.get(dk)
        if fn:
            if dk == "cover":
                fn(root, data, keyword, team, today)
            elif dk == "outro":
                fn(root, team, today)
            else:
                fn(root, dk, data, lbl())

        with open(path, "wb") as f:
            f.write(etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))

    with open(prs_path,  "wb") as f:
        f.write(etree.tostring(prs_root,  xml_declaration=True, encoding="UTF-8", standalone=True))
    with open(rels_path, "wb") as f:
        f.write(etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True))

    out_path = os.path.join(tmp_dir, f"{keyword}_발표자료.pptx")
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for rd, _, files in os.walk(unpacked):
            for file in files:
                fp = os.path.join(rd, file)
                zout.write(fp, os.path.relpath(fp, unpacked))

    background_tasks.add_task(shutil.rmtree, tmp_dir, ignore_errors=True)
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename=f"{keyword}_발표자료.pptx",
        background=background_tasks,
    )


# ──────────────────────────────────────────────
# 기타 엔드포인트
# ──────────────────────────────────────────────
@app.get("/")
def root():
    return {"status": "ok", "service": "모듀이 PPT 생성 API"}

@app.get("/api/meeting")
def meeting():
    return {"message": "회의록 AI 엔드포인트"}

@app.get("/api/output")
def output_api():
    return {"message": "산출물 AI 엔드포인트"}


# ──────────────────────────────────────────────
# Word 생성
# ──────────────────────────────────────────────
class WordRequest(BaseModel):
    keyword: str
    team: Optional[str] = "딸깍"
    model: Optional[str] = "gpt"

def call_gpt_word(keyword):
    client = openai.OpenAI(api_key=OPENAI_API_KEY)
    res = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {"role": "system", "content": """당신은 전문 문서 작성 전문가입니다. 키워드를 받아 아래 JSON 형식으로 Word 문서 내용을 작성하세요.

규칙:
1. JSON만 반환, 다른 텍스트 절대 금지
2. 모든 내용은 한국어(한글)로 작성
3. 한자 사용 절대 금지
4. 각 문단은 완결된 문장으로 작성
5. 전문적이고 격식체(~합니다, ~입니다)로 작성

{
  "title": "문서 제목 (20자이내)",
  "subtitle": "부제목 (30자이내)",
  "overview": {
    "heading": "1. 개요",
    "background": "배경 설명. 3~4문장. 150자이내",
    "purpose": "목적 설명. 2~3문장. 100자이내"
  },
  "main": {
    "heading": "2. 핵심 내용",
    "section1_title": "소제목1 (15자이내)",
    "section1_body": "소제목1 본문. 3~4문장. 150자이내",
    "section2_title": "소제목2 (15자이내)",
    "section2_body": "소제목2 본문. 3~4문장. 150자이내",
    "section3_title": "소제목3 (15자이내)",
    "section3_body": "소제목3 본문. 3~4문장. 150자이내"
  },
  "analysis": {
    "heading": "3. 분석",
    "cause1": "원인1. 2~3문장. 100자이내",
    "cause2": "원인2. 2~3문장. 100자이내",
    "cause3": "원인3. 2~3문장. 100자이내",
    "result": "분석 결과. 3~4문장. 150자이내"
  },
  "conclusion": {
    "heading": "4. 결론",
    "summary": "핵심 요약. 3~4문장. 150자이내",
    "expected": "기대효과. 2~3문장. 100자이내",
    "one_line_summary": "이 보고서 전체를 한 문장으로 요약. 50자이내"
  }
}"""},
            {"role": "user", "content": f"키워드: {keyword}"}
        ]
    )
    clean = re.sub(r"```json|```", "", res.choices[0].message.content).strip()
    return json.loads(clean)

def create_word(keyword, team, today, data):
    import copy
    from docx import Document
    from docx.oxml.ns import qn

    TEMPLATE_PATH_WORD = os.path.join(os.path.dirname(__file__), "template_word.docx")
    doc = Document(TEMPLATE_PATH_WORD)

    W   = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"

    def get_txbx_text(txbx):
        return ''.join(t.text or '' for t in txbx.iter(f"{{{W}}}t")).strip()

    for txbx in doc.element.body.iter(f"{{{WPS}}}txbx"):
        text = get_txbx_text(txbx)
        paras_tb = list(txbx.iter(f"{{{W}}}p"))
        runs_with_text = [p for p in paras_tb if list(p.iter(f"{{{W}}}r"))]

        if 'Letter of Intent' in text or 'Purchase' in text:
            for t in txbx.iter(f"{{{W}}}t"):
                t.text = ''
            if len(runs_with_text) >= 2:
                ts0 = list(runs_with_text[0].iter(f"{{{W}}}t"))
                if ts0: ts0[0].text = data["title"][:20]
                ts1 = list(runs_with_text[1].iter(f"{{{W}}}t"))
                if ts1: ts1[0].text = data["subtitle"][:30]
            elif len(runs_with_text) == 1:
                ts0 = list(runs_with_text[0].iter(f"{{{W}}}t"))
                if ts0: ts0[0].text = data["title"][:20]

        elif 'WEBERSTEIN' in text.upper():
            for t in txbx.iter(f"{{{W}}}t"):
                t.text = team.upper()
                break

    def mk_run_br(para, text, bold=False, base_rPr=None):
        lines = text.split('\n')
        for i, line in enumerate(lines):
            r = etree.SubElement(para._element, f"{{{W}}}r")
            if base_rPr is not None:
                rp = copy.deepcopy(base_rPr)
                b_el = rp.find(f"{{{W}}}b")
                if bold and b_el is None:
                    etree.SubElement(rp, f"{{{W}}}b")
                elif not bold and b_el is not None:
                    rp.remove(b_el)
                r.insert(0, rp)
            t = etree.SubElement(r, f"{{{W}}}t")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t.text = line
            if i < len(lines) - 1:
                etree.SubElement(r, f"{{{W}}}br")

    def get_base_rPr(para):
        for run in para.runs:
            rPr = run._element.find(f"{{{W}}}rPr")
            if rPr is not None:
                fonts = rPr.find(f"{{{W}}}rFonts")
                if fonts is not None:
                    font = fonts.get(f"{{{W}}}ascii") or fonts.get(f"{{{W}}}eastAsia") or ""
                    if "MS Gothic" not in font:
                        return rPr
                else:
                    return rPr
        return None

    def clear_runs(para):
        for r in list(para._element.findall(f"{{{W}}}r")):
            para._element.remove(r)

    paras = doc.paragraphs

    if len(paras) > 6 and paras[6].runs:
        paras[6].runs[0].text = today

    if len(paras) > 7:
        para = paras[7]
        rPr = get_base_rPr(para)
        clear_runs(para)
        mk_run_br(para, data["overview"]["background"], base_rPr=rPr)

    if len(paras) > 10:
        para = paras[10]
        rPr = get_base_rPr(para)
        clear_runs(para)
        mk_run_br(para, data["overview"]["purpose"], base_rPr=rPr)

    if len(paras) > 11:
        main = data["main"]
        para = paras[11]
        rPr = get_base_rPr(para)
        clear_runs(para)
        mk_run_br(para, main['section1_title'] + "\n", bold=True, base_rPr=rPr)
        mk_run_br(para, main['section1_body'] + "\n\n", base_rPr=rPr)
        mk_run_br(para, main['section2_title'] + "\n", bold=True, base_rPr=rPr)
        mk_run_br(para, main['section2_body'] + "\n\n", base_rPr=rPr)
        mk_run_br(para, main['section3_title'] + "\n", bold=True, base_rPr=rPr)
        mk_run_br(para, main['section3_body'], base_rPr=rPr)

    if len(paras) > 12:
        analysis   = data["analysis"]
        conclusion = data["conclusion"]
        para = paras[12]
        rPr = get_base_rPr(para)
        clear_runs(para)
        mk_run_br(para, "\n분석\n", bold=True, base_rPr=rPr)
        mk_run_br(para, "• " + analysis['cause1'] + "\n", base_rPr=rPr)
        mk_run_br(para, "• " + analysis['cause2'] + "\n", base_rPr=rPr)
        mk_run_br(para, "• " + analysis['cause3'] + "\n", base_rPr=rPr)
        mk_run_br(para, analysis['result'] + "\n\n", base_rPr=rPr)
        mk_run_br(para, "결론\n", bold=True, base_rPr=rPr)
        mk_run_br(para, conclusion['summary'] + "\n", base_rPr=rPr)
        mk_run_br(para, conclusion['expected'], base_rPr=rPr)

    if len(paras) > 13:
        para = paras[13]
        rPr = get_base_rPr(para)
        clear_runs(para)
        one_line = data.get("conclusion", {}).get("one_line_summary", "")
        mk_run_br(para, "\n핵심 요약\n", bold=True, base_rPr=rPr)
        mk_run_br(para, one_line + "\n\n", base_rPr=rPr)
        mk_run_br(para, f"팀  {team}  |  {today}  |  modui.ai", base_rPr=rPr)

    return doc

@app.post("/generate-word")
def generate_word(req: WordRequest, background_tasks: BackgroundTasks):
    keyword = req.keyword
    team    = req.team or "딸깍"
    today   = date.today().strftime("%Y.%m.%d")

    data = call_gemini_word(keyword) if req.model == "gemini" else call_gpt_word(keyword)
    data = sanitize_data(data)

    tmp_dir  = tempfile.mkdtemp()
    out_path = os.path.join(tmp_dir, f"{keyword}_문서.docx")

    doc = create_word(keyword, team, today, data)
    doc.save(out_path)

    background_tasks.add_task(shutil.rmtree, tmp_dir, ignore_errors=True)
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"{keyword}_문서.docx",
        background=background_tasks,
    )


# ──────────────────────────────────────────────
# PDF 생성 (Word → LibreOffice 변환)
# ──────────────────────────────────────────────
class PDFRequest(BaseModel):
    keyword: str
    team: Optional[str] = "딸깍"
    model: Optional[str] = "gpt"

@app.post("/generate-pdf")
def generate_pdf(req: PDFRequest, background_tasks: BackgroundTasks):
    keyword = req.keyword
    team    = req.team or "딸깍"
    today   = date.today().strftime("%Y.%m.%d")

    data = call_gemini_word(keyword) if req.model == "gemini" else call_gpt_word(keyword)
    data = sanitize_data(data)

    tmp_dir   = tempfile.mkdtemp()
    docx_path = os.path.join(tmp_dir, f"{keyword}_문서.docx")
    pdf_path  = os.path.join(tmp_dir, f"{keyword}_문서.pdf")

    doc = create_word(keyword, team, today, data)
    doc.save(docx_path)

    subprocess.run([
        "libreoffice", "--headless", "--convert-to", "pdf",
        "--outdir", tmp_dir, docx_path
    ], check=True)

    background_tasks.add_task(shutil.rmtree, tmp_dir, ignore_errors=True)
    return FileResponse(
        pdf_path,
        media_type="application/pdf",
        filename=f"{keyword}_문서.pdf",
        background=background_tasks,
    )


# ──────────────────────────────────────────────
# 회의록 AI
# ──────────────────────────────────────────────
@app.post("/analyze-minutes")
async def analyze_minutes(
    file: UploadFile = File(...),
    title: str = Form("회의"),
    date: str = Form(""),
    members: str = Form(""),
    model: str = Form("gpt"),
):
    tmp_dir = tempfile.mkdtemp()
    try:
        audio_path = os.path.join(tmp_dir, file.filename)
        with open(audio_path, "wb") as f:
            f.write(await file.read())

        client_stt = openai.OpenAI(api_key=OPENAI_API_KEY)
        with open(audio_path, "rb") as audio_file:
            transcript = client_stt.audio.transcriptions.create(
                model="whisper-1",
                file=audio_file,
                language="ko"
            ).text

        if model == "gemini":
            summary = call_gemini_minutes(transcript, title, members)
            summary = sanitize_data(summary)
            return {"title": title, "date": date, "members": members, "transcript": transcript, **summary}

        client = openai.OpenAI(api_key=OPENAI_API_KEY)
        res = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": """당신은 회의록 작성 전문가입니다. 회의 내용을 분석하여 아래 JSON 형식으로 반환하세요.

규칙:
1. JSON만 반환, 다른 텍스트 절대 금지
2. 모든 내용은 한국어로 작성
3. 한자 사용 절대 금지

{
  "agenda": [
    {"title": "안건 제목", "content": "안건 내용 요약"}
  ],
  "summary": "회의 전체 내용을 3줄로 요약",
  "action_items": [
    {"content": "해야 할 일", "deadline": "기한 (없으면 '미정')"}
  ],
  "next_agenda": "다음 회의에서 논의할 안건"
}"""},
                {"role": "user", "content": f"회의 제목: {title}\n참석자: {members}\n\n회의 내용:\n{transcript}"}
            ]
        )
        clean = re.sub(r"```json|```", "", res.choices[0].message.content).strip()
        summary = json.loads(clean)
        summary = sanitize_data(summary)

        return {
            "title": title,
            "date": date,
            "members": members,
            "transcript": transcript,
            **summary
        }

    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
# ──────────────────────────────────────────────
# 회의록 → Word 공문 생성
# ──────────────────────────────────────────────
class MinutesDocRequest(BaseModel):
    title: str
    date: Optional[str] = ""
    members: Optional[str] = ""
    transcript: Optional[str] = ""
    agenda: Optional[list] = []
    summary: Optional[str] = ""
    action_items: Optional[list] = []
    next_agenda: Optional[str] = ""

@app.post("/generate-minutes-word")
def generate_minutes_word(req: MinutesDocRequest, background_tasks: BackgroundTasks):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import copy

    today = req.date or date.today().strftime("%Y.%m.%d")
    doc_num = f"MOD-{today.replace('.','').replace('-','')[:8]}-001"

    doc = Document()

    # 여백 설정
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(3.0)

    COLOR_PRIMARY = RGBColor(0x1B, 0x3A, 0x6B)
    COLOR_ACCENT  = RGBColor(0x2E, 0x5F, 0xA3)
    COLOR_GRAY    = RGBColor(0x66, 0x66, 0x66)
    FONT_NAME     = "맑은 고딕"

    def add_run(para, text, bold=False, size=10, color=None):
        run = para.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        return run

    def set_cell_bg(cell, hex_color):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

    def add_section_title(label):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(28)
        p.paragraph_format.space_after  = Pt(8)
        # 왼쪽 파란 border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '18')
        left.set(qn('w:space'), '8')
        left.set(qn('w:color'), '2E5FA3')
        pBdr.append(left)
        pPr.append(pBdr)
        p.paragraph_format.left_indent = Cm(0.4)
        add_run(p, label, bold=True, size=12, color=COLOR_PRIMARY)

    # ── 제목
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    add_run(title_p, "회  의  록", bold=True, size=20, color=COLOR_PRIMARY)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(2)
    add_run(sub_p, req.title, bold=True, size=12, color=COLOR_ACCENT)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.space_after = Pt(14)
    add_run(date_p, f"{today}  |  modui.ai", size=9, color=COLOR_GRAY)

    doc.add_paragraph()  # 구분 공백

    # ── 1. 회의 개요 테이블
    add_section_title("1. 회의 개요")
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = 'Table Grid'
    labels = ["문서번호", "문서명", "회의명", "일  시", "참 석 자"]
    values = [doc_num, "회의록", req.title, today, req.members]
    for i, (lbl, val) in enumerate(zip(labels, values)):
        row = info_table.rows[i]
        row.cells[0].width = Cm(3.0)
        row.cells[1].width = Cm(12.0)
        set_cell_bg(row.cells[0], 'E8EEF7')
        lp = row.cells[0].paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(lp, lbl, bold=True, size=9, color=COLOR_PRIMARY)
        vp = row.cells[1].paragraphs[0]
        add_run(vp, val or "-", size=9)
    doc.add_paragraph()
    doc.add_paragraph()

    # ── 2. 안건
    add_section_title("2. 안건 및 논의 내용")
    ag_table = doc.add_table(rows=1 + len(req.agenda), cols=3)
    ag_table.style = 'Table Grid'
    # 헤더
    hrow = ag_table.rows[0]
    for cell, label, w in zip(hrow.cells, ["번호","안건","내용"], [Cm(1.5),Cm(4.0),Cm(9.5)]):
        cell.width = w
        set_cell_bg(cell, '1B3A6B')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, label, bold=True, size=9, color=RGBColor(0xFF,0xFF,0xFF))
    # 데이터
    for i, item in enumerate(req.agenda or []):
        row = ag_table.rows[i+1]
        bg = 'FFFFFF' if i % 2 == 0 else 'E8EEF7'
        for cell in row.cells:
            set_cell_bg(cell, bg)
        row.cells[0].width = Cm(1.5)
        row.cells[1].width = Cm(4.0)
        row.cells[2].width = Cm(9.5)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p0, str(i+1), bold=True, size=9, color=COLOR_ACCENT)
        add_run(row.cells[1].paragraphs[0], item.get("title",""), bold=True, size=9)
        add_run(row.cells[2].paragraphs[0], item.get("content",""), size=9)
    doc.add_paragraph()
    doc.add_paragraph()

    # ── 3. 회의 요약
    add_section_title("3. 회의 요약")
    for line in (req.summary or "").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(3)
        add_run(p, "• ", bold=True, color=COLOR_ACCENT)
        add_run(p, line.strip(), size=9)
    doc.add_paragraph()
    doc.add_paragraph()

    # ── 4. Action Items
    add_section_title("4. 결정사항 및 후속 조치")
    ai_table = doc.add_table(rows=1 + len(req.action_items), cols=3)
    ai_table.style = 'Table Grid'
    hrow2 = ai_table.rows[0]
    for cell, label, w in zip(hrow2.cells, ["No.","할 일","기한"], [Cm(1.2),Cm(10.3),Cm(3.5)]):
        cell.width = w
        set_cell_bg(cell, '2E5FA3')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, label, bold=True, size=9, color=RGBColor(0xFF,0xFF,0xFF))
    for i, item in enumerate(req.action_items or []):
        row = ai_table.rows[i+1]
        bg = 'FFFFFF' if i % 2 == 0 else 'E8EEF7'
        for cell in row.cells:
            set_cell_bg(cell, bg)
        row.cells[0].width = Cm(1.2)
        row.cells[1].width = Cm(10.3)
        row.cells[2].width = Cm(3.5)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p0, str(i+1), bold=True, size=9, color=COLOR_ACCENT)
        add_run(row.cells[1].paragraphs[0], item.get("content",""), size=9)
        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        deadline = item.get("deadline","미정")
        add_run(p2, deadline, bold=(deadline != "미정"), size=9,
                color=COLOR_GRAY if deadline == "미정" else COLOR_PRIMARY)
    doc.add_paragraph()

    # ── 5. 차기 안건
    add_section_title("5. 차기 회의 안건")
    next_p = doc.add_paragraph()
    next_p.paragraph_format.left_indent = Cm(0.4)
    add_run(next_p, req.next_agenda or "-", size=9)

    # ── 저장
    tmp_dir  = tempfile.mkdtemp()
    filename = f"{req.title}_회의록.docx"
    out_path = os.path.join(tmp_dir, filename)
    doc.save(out_path)

    background_tasks.add_task(shutil.rmtree, tmp_dir, ignore_errors=True)
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename,
        background=background_tasks,
    )