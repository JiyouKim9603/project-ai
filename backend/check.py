from pptx import Presentation
prs = Presentation("template.pptx")
slide8 = prs.slides[7]
for shape in slide8.shapes:
    if shape.name in ["TextBox 33","TextBox 34","TextBox 35"]:
        print(f"{shape.name} | w={shape.width/914400*2.54:.1f}cm | h={shape.height/914400*2.54:.1f}cm")